from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.core.config import settings
from app.extraction_review import build_extraction_review
from app.quote_engine import calculate_quote
from app.quote_workflow import customer_export_readiness
from tests.test_web_quote_new import _logged_in_client


SPEC = {"layer": 6, "qty": 10, "material": "FR4", "length_mm": 100, "width_mm": 80,
        "thickness_mm": 1.6, "copper_weight": "1oz", "surface_finish": "OSP",
        "delivery_days": 7, "pitch_mm": 0.4}


def quote(spec=None, status="approved"):
    spec = deepcopy(SPEC if spec is None else spec)
    result = calculate_quote(spec)
    return SimpleNamespace(spec_json=spec, breakdown_json=result, total=result.get("total"),
                           unit_price=result.get("unit_price"), currency="NTD", status=status)


def codes(readiness):
    return {b["code"] for b in readiness["formal_blockers"]}


def saved_quote(db, value):
    assert db.save_quote("web:1", value.spec_json, value.breakdown_json)
    with db.SessionLocal() as session:
        saved = session.query(db.QuoteHistory).order_by(db.QuoteHistory.id.desc()).first()
        saved.status = value.status
        saved.currency = value.currency
        session.commit()
        return saved.id


def test_complete_approved_quote_and_ordered_quote_are_ready():
    for status in ("approved", "ordered"):
        value = quote(status=status)
        before = deepcopy(vars(value))
        assert customer_export_readiness(value)["formal_ready"]
        assert vars(value) == before


def test_review_and_business_approval_are_separate_gates():
    value = quote(status="pending")
    readiness = customer_export_readiness(value)
    assert readiness["estimate_ready"] and not readiness["formal_ready"]
    assert codes(readiness) == {"export_approval_required"}
    value.spec_json["_extraction_review"] = build_extraction_review(value.spec_json, "6L qty 10")
    value.status = "approved"
    readiness = customer_export_readiness(value)
    assert not readiness["estimate_ready"] and "export_pending_review" in codes(readiness)


@pytest.mark.parametrize("field", ["material", "surface_finish", "copper_weight", "thickness_mm", "delivery_days", "pitch_mm"])
def test_missing_specs_still_allow_only_an_estimate(field):
    spec = deepcopy(SPEC)
    del spec[field]
    readiness = customer_export_readiness(quote(spec))
    assert readiness["estimate_ready"] and not readiness["formal_ready"]
    assert "export_missing_specs" in codes(readiness)


@pytest.mark.parametrize("field", ["pitch_mm", "delivery_days", "thickness_mm", "copper_weight"])
@pytest.mark.parametrize("value", [0, -1, "nan", "inf"])
def test_invalid_numeric_specs_cannot_be_formal(field, value):
    item = quote()
    item.spec_json[field] = value
    assert not customer_export_readiness(item)["formal_ready"]


def test_split_copper_and_area_only_specs_are_supported():
    spec = deepcopy(SPEC)
    for field in ("copper_weight", "length_mm", "width_mm"):
        del spec[field]
    spec.update(copper_outer_oz=2, copper_inner_oz=1, area_inch=12)
    assert customer_export_readiness(quote(spec))["formal_ready"]


@pytest.mark.parametrize("finish", ["ENIG", "Hard Gold"])
def test_gold_finish_requires_explicit_positive_thickness(finish):
    spec = {**SPEC, "surface_finish": finish, "enig": finish == "ENIG", "hard_gold": finish == "Hard Gold"}
    assert not customer_export_readiness(quote(spec))["formal_ready"]
    spec["enig_thickness_uinch"] = 5
    assert customer_export_readiness(quote(spec))["formal_ready"]


@pytest.mark.parametrize("finish,enig,hard_gold", [("ENIG", False, False), ("ENIG", True, True),
    ("Hard Gold", False, False), ("OSP", True, False), ("HASL", False, True)])
def test_finish_and_priced_process_flags_must_agree(finish, enig, hard_gold):
    spec = {**SPEC, "surface_finish": finish, "enig": enig, "hard_gold": hard_gold, "enig_thickness_uinch": 5}
    readiness = customer_export_readiness(quote(spec))
    assert readiness["estimate_ready"] and not readiness["formal_ready"]
    assert "export_process_conflict" in codes(readiness)


def test_unknown_finish_is_not_a_complete_specification():
    item = quote()
    item.spec_json["surface_finish"] = "TBD"
    assert not customer_export_readiness(item)["formal_ready"]


def test_real_unpriced_inspection_request_cannot_be_released_by_approval():
    spec = deepcopy(SPEC)
    spec["inspection_report_required"] = True
    readiness = customer_export_readiness(quote(spec))
    assert readiness["estimate_ready"]
    assert "export_pricing_not_ready" in codes(readiness)


@pytest.mark.parametrize("pricing", [None, {}, {"status": "quotable"},
    {"status": "quotable", "critical_missing": [], "unpriced_factors": ["Inspection"]},
    {"status": "estimate", "critical_missing": [], "unpriced_factors": []},
    {"status": "needs_review", "critical_missing": [], "unpriced_factors": []}])
def test_missing_or_unresolved_pricing_review_fails_closed(pricing):
    item = quote()
    item.breakdown_json["pricing_review"] = pricing
    readiness = customer_export_readiness(item)
    assert readiness["estimate_ready"] and not readiness["formal_ready"]


@pytest.mark.parametrize("field,value", [("currency", "UNKNOWN"), ("total", float("nan")),
    ("total", 123), ("unit_price", -1), ("unit_price", float("inf"))])
def test_invalid_saved_financial_data_blocks_both_documents(field, value):
    item = quote()
    setattr(item, field, value)
    readiness = customer_export_readiness(item)
    assert not readiness["formal_ready"] and not readiness["estimate_ready"]


def test_missing_core_specs_block_customer_estimates():
    item = quote()
    item.spec_json["qty"] = 0
    assert not customer_export_readiness(item)["estimate_ready"]


def test_incomplete_approved_quote_endpoint_and_workbook(temp_db, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "EXPORT_DIR", str(tmp_path))
    client = _logged_in_client(temp_db)
    spec = deepcopy(SPEC)
    del spec["surface_finish"]
    del spec["pitch_mm"]
    value = quote(spec)
    qid = saved_quote(temp_db, value)
    blocked = client.get(f"/quotes/{qid}/export/formal")
    assert blocked.status_code == 409
    assert "export_missing_specs" in {b["code"] for b in blocked.json()["detail"]["blockers"]}
    assert list(tmp_path.iterdir()) == []
    html = client.get(f"/quotes/{qid}").text
    assert f'href="/quotes/{qid}/export/formal"' not in html
    assert f'href="/quotes/{qid}/export/estimate"' in html
    assert "Surface Finish, Pitch" in html
    download = client.get(f"/quotes/{qid}/export/estimate", follow_redirects=False)
    assert download.status_code == 303
    assert "/estimate_" in download.headers["location"]
    content = client.get(download.headers["location"]).content
    workbook = load_workbook(BytesIO(content))
    values = [cell.value for row in workbook.active for cell in row if cell.value]
    assert workbook.active["A3"].value == "PRELIMINARY ESTIMATE - NOT AN OFFICIAL QUOTATION"
    assert workbook.active["E11"].value == f"NTD {value.total:,.2f}"
    assert "Missing specifications: Surface Finish, Pitch" in values
    assert not any("NET 30" in str(v) or "Validity: 30 days" in str(v) for v in values)
    workbook.close()
    with temp_db.SessionLocal() as session:
        saved = session.get(temp_db.QuoteHistory, qid)
        assert saved.spec_json == value.spec_json and saved.total == value.total


def test_complete_quote_requires_approval_and_pending_estimate_is_labelled(temp_db, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "EXPORT_DIR", str(tmp_path))
    client = _logged_in_client(temp_db)
    qid = saved_quote(temp_db, quote(status="pending"))
    assert client.get(f"/quotes/{qid}/export/formal").status_code == 409
    estimate = client.get(f"/quotes/{qid}/export/estimate", follow_redirects=False)
    assert estimate.status_code == 303
    assert client.patch(f"/api/quotes/{qid}", json={"status": "approved"}).status_code == 200
    formal = client.get(f"/quotes/{qid}/export/formal", follow_redirects=False)
    assert formal.status_code == 303
    workbook = load_workbook(BytesIO(client.get(formal.headers["location"]).content))
    assert "Official Quotation" in workbook.active["A3"].value
    workbook.close()
    assert client.patch(f"/api/quotes/{qid}", json={"status": "pending"}).status_code == 200
    assert client.get(f"/quotes/{qid}/export/formal").status_code == 409


def test_pending_review_blocks_estimate_and_formal_without_writing_files(temp_db, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "EXPORT_DIR", str(tmp_path))
    client = _logged_in_client(temp_db)
    value = quote()
    value.spec_json["_extraction_review"] = build_extraction_review(value.spec_json, "6L qty 10")
    qid = saved_quote(temp_db, value)
    for kind in ("estimate", "formal"):
        assert client.get(f"/quotes/{qid}/export/{kind}").status_code == 409
    assert list(tmp_path.iterdir()) == []
    client.cookies.clear()
    redirect = client.get(f"/quotes/{qid}/export/estimate", follow_redirects=False)
    assert redirect.status_code == 303 and redirect.headers["location"] == "/login"


def test_line_formal_command_cannot_bypass_staff_release(temp_db, monkeypatch):
    import app.main as main

    replies = []
    def forbidden(*args, **kwargs):
        raise AssertionError("LINE formal command must not parse, save or export a quote")
    monkeypatch.setattr(main, "export_formal_quote", forbidden)
    monkeypatch.setattr(main, "calculate_quote", forbidden)
    monkeypatch.setattr(main, "save_quote", forbidden)
    monkeypatch.setattr(main, "MessagingApi", lambda client: SimpleNamespace(reply_message=replies.append))
    event = SimpleNamespace(source=SimpleNamespace(user_id="synthetic-line-user"),
                            message=SimpleNamespace(text="Formal Quote"), reply_token="synthetic-reply")
    main.handle_message(event)
    assert len(replies) == 1
    assert "Formal quotes require staff review" in replies[0].messages[0].text
