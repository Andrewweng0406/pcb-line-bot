import os
import socket
import threading
import time
from copy import deepcopy
from pathlib import Path

import pytest

from app.core.config import settings
from app.extraction_review import release_pending
from tests.test_quote_acceptance import RFQ, PARSED_RFQ, workbook_values
from tests.test_web_quote_new import _logged_in_client


pytestmark = pytest.mark.skipif(
    os.environ.get("RUN_BROWSER_ACCEPTANCE") != "1",
    reason="Set RUN_BROWSER_ACCEPTANCE=1 with Playwright Chromium installed.",
)


@pytest.fixture
def acceptance_server(temp_db, monkeypatch, tmp_path):
    import uvicorn
    import app.web as web
    from app.main import app

    monkeypatch.setattr(settings, "EXPORT_DIR", str(tmp_path))
    monkeypatch.setattr(settings, "DEFAULT_CURRENCY", "NTD")
    monkeypatch.setattr(settings, "PRICING_VERSION", "browser-acceptance-v1")
    monkeypatch.setattr(web, "parse_pcb_text", lambda text: deepcopy(PARSED_RFQ))
    _logged_in_client(temp_db).close()
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        listener.listen()
        port = listener.getsockname()[1]
        server = uvicorn.Server(uvicorn.Config(app, log_level="warning"))
        thread = threading.Thread(target=server.run, kwargs={"sockets": [listener]}, daemon=True)
        thread.start()
        try:
            deadline = time.monotonic() + 10
            while not server.started and thread.is_alive() and time.monotonic() < deadline:
                time.sleep(0.01)
            assert server.started, "Acceptance server failed to start"
            yield f"http://127.0.0.1:{port}", temp_db
        finally:
            server.should_exit = True
            thread.join(timeout=10)
            assert not thread.is_alive(), "Acceptance server did not stop"


def confirm_fields(page, note):
    form = page.locator("#extraction-review form")
    fields = form.locator('input[name="reviewed_fields"]')
    assert fields.count() > 0
    for field in fields.all():
        field.check()
    form.locator('textarea[name="review_note"]').fill(note)
    with page.expect_navigation():
        form.get_by_role("button", name="Confirm Selected Fields", exact=True).click()


def approve(page):
    page.get_by_text("Status and Internal Notes", exact=True).click()
    form = page.locator('form[action$="/update"]').first
    form.locator('select[name="status"]').select_option("approved")
    with page.expect_navigation():
        form.get_by_role("button", name="Save", exact=True).click()


@pytest.mark.parametrize("width", [390, 1440])
@pytest.mark.parametrize("role", ["viewer", "staff", "manager"])
def test_browser_role_controls(acceptance_server, width, role, tmp_path):
    from playwright.sync_api import sync_playwright
    from app.core.auth import hash_password

    base, db = acceptance_server
    with db.SessionLocal() as session:
        session.add(db.User(email=f"role-{role}@example.com", password_hash=hash_password("test-password"), role=role))
        session.commit()
    from app.quote_engine import calculate_quote
    spec = {"layer": 6, "qty": 10, "area_inch": 10, "material": "FR4", "surface_finish": "OSP",
            "copper_weight": "1oz", "thickness_mm": 1.6, "pitch_mm": 0.4, "delivery_days": 7}
    db.save_quote("web:1", spec, calculate_quote(spec))
    with db.SessionLocal() as session:
        quote = session.query(db.QuoteHistory).first()
        qid = quote.id
        if role == "manager":
            quote.status = "approved"
            session.commit()
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        page = browser.new_page(viewport={"width": width, "height": 900})
        errors = []
        page.on("pageerror", lambda error: errors.append(str(error)))
        page.goto(base + "/login")
        page.locator('input[name="email"]').fill(f"role-{role}@example.com")
        page.locator('input[name="password"]').fill("test-password")
        page.get_by_role("button", name="Log In", exact=True).click()
        page.wait_for_url(base + "/")
        page.goto(base + f"/quotes/{qid}")
        assert page.locator('select[name="status"]').count() == (1 if role == "manager" else 0)
        assert page.locator('input[name="final_price"]').count() == (1 if role == "manager" else 0)
        assert page.locator('a[href$="/export/formal"]').count() == (1 if role == "manager" else 0)
        assert page.locator(f'a[href="/quotes/{qid}/revise"]').count() == (0 if role == "viewer" else 1)
        assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
        if role == "staff":
            page.get_by_text("Status and Internal Notes", exact=True).click()
            form = page.locator(f'form[action="/quotes/{qid}/update"]').first
            form.locator('textarea[name="notes"]').fill("Staff review note")
            form.get_by_role("button", name="Save", exact=True).click()
            with db.SessionLocal() as session:
                quote = session.get(db.QuoteHistory, qid)
                assert quote.notes == "Staff review note" and quote.status == "pending"
        page.screenshot(path=str(tmp_path / f"role-{role}-{width}.png"))
        assert not errors
        browser.close()


@pytest.mark.parametrize("width", [390, 1440])
def test_browser_import_preview_preserves_mapping_and_layout(acceptance_server, width, tmp_path):
    from io import BytesIO
    from openpyxl import Workbook
    from playwright.sync_api import sync_playwright

    base, db = acceptance_server
    workbook = Workbook()
    workbook.active.append(["Layers", "Qty", "Quote", "Billing Currency"])
    workbook.active.append([6, 10, 1000, "USD"])
    workbook.active.append([6, 0, 1000, "USD"])
    output = BytesIO()
    workbook.save(output)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        page = browser.new_page(viewport={"width": width, "height": 900})
        page.goto(base + "/login")
        page.locator('input[name="email"]').fill("staff@example.com")
        page.locator('input[name="password"]').fill("hunter2")
        page.get_by_role("button", name="Log In", exact=True).click()
        page.wait_for_url(base + "/")
        page.goto(base + "/import/quotes")
        page.locator("details summary").click()
        page.locator('input[name="currency_col"]').fill("Billing Currency")
        page.locator('input[type="file"]').set_input_files({"name": "synthetic.xlsx", "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "buffer": output.getvalue()})
        page.get_by_role("button", name="Preview / Dry Run", exact=True).click()
        heading = page.get_by_role("heading", name="Import Preview", exact=True)
        heading.wait_for()
        assert page.locator('input[name="currency_col"]').input_value() == "Billing Currency"
        heading.scroll_into_view_if_needed()
        page.get_by_role("region", name="Import rows").evaluate("e => e.scrollLeft = 100")
        box = heading.bounding_box()
        assert box["x"] >= 0 and box["x"] + box["width"] <= width
        assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
        page.screenshot(path=str(tmp_path / f"import-preview-{width}.png"))
        with db.SessionLocal() as session:
            assert session.query(db.QuoteHistory).count() == 0
        browser.close()


@pytest.mark.parametrize("width", [390, 1440])
def test_browser_incomplete_approved_quote_only_exports_estimate(acceptance_server, width, tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook
    from playwright.sync_api import sync_playwright
    from app.quote_engine import calculate_quote

    base, db = acceptance_server
    spec = {"layer": 6, "qty": 10, "length_mm": 100, "width_mm": 80, "material": "FR4",
            "thickness_mm": 1.6, "copper_weight": "1oz", "delivery_days": 7}
    assert db.save_quote("web:1", spec, calculate_quote(spec))
    with db.SessionLocal() as session:
        quote = session.query(db.QuoteHistory).one()
        qid = quote.id
        quote.status = "approved"
        session.commit()
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        page = browser.new_page(viewport={"width": width, "height": 900}, accept_downloads=True)
        errors = []
        page.on("pageerror", lambda error: errors.append(str(error)))
        page.goto(base + "/login")
        page.locator('input[name="email"]').fill("staff@example.com")
        page.locator('input[name="password"]').fill("hunter2")
        page.get_by_role("button", name="Log In", exact=True).click()
        page.wait_for_url(base + "/")
        page.goto(base + f"/quotes/{qid}")
        assert page.get_by_text("Formal export blocked", exact=True).is_visible()
        assert page.get_by_role("link", name="Generate Formal Quote", exact=True).count() == 0
        assert page.request.get(base + f"/quotes/{qid}/export/formal").status == 409
        with page.expect_download() as event:
            page.get_by_role("link", name="Download Estimate", exact=True).click()
        assert event.value.failure() is None
        workbook = load_workbook(BytesIO(Path(event.value.path()).read_bytes()))
        assert workbook.active["A3"].value == "PRELIMINARY ESTIMATE - NOT AN OFFICIAL QUOTATION"
        workbook.close()
        assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
        page.screenshot(path=str(tmp_path / f"estimate-release-{width}.png"), full_page=True)
        assert not errors
        with db.SessionLocal() as session:
            quote = session.get(db.QuoteHistory, qid)
            assert quote.status == "approved" and quote.spec_json == spec
        browser.close()


@pytest.mark.parametrize("width", [390, 1440])
def test_browser_rfq_review_export_and_revision(acceptance_server, width, tmp_path):
    from playwright.sync_api import sync_playwright

    base, db = acceptance_server
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        context = browser.new_context(viewport={"width": width, "height": 900}, accept_downloads=True)
        page = context.new_page()
        errors = []
        page.on("pageerror", lambda error: errors.append(str(error)))
        page.goto(base + "/login")
        page.locator('input[name="email"]').fill("staff@example.com")
        page.locator('input[name="password"]').fill("hunter2")
        page.get_by_role("button", name="Log In", exact=True).click()
        page.wait_for_url(base + "/")
        page.goto(base + "/quotes/new")
        page.locator("#ai-spec-text").fill(RFQ)
        page.get_by_role("button", name="Parse with AI", exact=True).click()
        page.locator('input[name="extraction_review_token"]').wait_for(state="attached")
        page.locator('input[name="layer"]').fill("8")
        page.locator('input[name="company_name"]').fill("Browser Acceptance Buyer")
        page.get_by_role("button", name="Calculate and Save Quote", exact=True).click()
        page.wait_for_url(base + "/quotes")
        with db.SessionLocal() as session:
            parent = session.query(db.QuoteHistory).one()
            parent_id = parent.id
            assert parent.layer == 8 and release_pending(parent.spec_json)
        page.goto(f"{base}/quotes/{parent_id}")
        assert page.request.get(f"{base}/quotes/{parent_id}/export/formal").status == 409
        assert page.get_by_role("link", name="Generate Formal Quote", exact=True).count() == 0
        page.get_by_role("link", name="Review Pending Fields", exact=True).click()
        confirm_fields(page, "Customer confirmed 8 layers, drawing Rev B")
        approve(page)
        assert page.locator("[data-price-status]").get_attribute("data-price-status") == "insufficient_evidence"
        with db.SessionLocal() as session:
            parent = session.get(db.QuoteHistory, parent_id)
            snapshot = deepcopy(parent.spec_json)
            assert parent.status == "approved" and not release_pending(snapshot)
        with page.expect_download() as event:
            page.get_by_role("link", name="Generate Formal Quote", exact=True).click()
        download = event.value
        assert download.failure() is None
        original_bytes = Path(download.path()).read_bytes()
        assert "Layer: 8" in workbook_values(original_bytes)["B11"]
        original_url = base + "/download/exports/" + download.suggested_filename
        page.get_by_role("link", name="Create Revision", exact=True).click()
        page.wait_for_url(f"{base}/quotes/{parent_id}/revise")
        page.locator('input[name="layer"]').fill("10")
        page.get_by_role("button", name="Calculate and Save Quote", exact=True).click()
        page.wait_for_url(base + "/quotes")
        with db.SessionLocal() as session:
            child = session.query(db.QuoteHistory).order_by(db.QuoteHistory.id.desc()).first()
            child_id = child.id
            assert child.layer == 10 and child.status == "pending"
            assert child.spec_json["_extraction_review"]["revision_of"] == parent_id
            assert "layer" in [item["field"] for item in release_pending(child.spec_json)]
        page.goto(f"{base}/quotes/{child_id}")
        assert page.request.get(f"{base}/quotes/{child_id}/export/formal").status == 409
        confirm_fields(page, "Customer requested 10 layers, drawing Rev C")
        approve(page)
        assert page.locator("[data-price-status]").get_attribute("data-price-status") == "insufficient_evidence"
        with page.expect_download() as event:
            page.get_by_role("link", name="Generate Formal Quote", exact=True).click()
        assert "Layer: 10" in workbook_values(Path(event.value.path()).read_bytes())["B11"]
        assert page.request.get(original_url).body() == original_bytes
        assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
        page.screenshot(path=str(tmp_path / f"acceptance-{width}.png"), full_page=True)
        with db.SessionLocal() as session:
            parent = session.get(db.QuoteHistory, parent_id)
            child = session.get(db.QuoteHistory, child_id)
            assert parent.spec_json == snapshot
            assert child.spec_json["_extraction_review"]["raw_input"] == RFQ
            assert any(event["note"].endswith("Rev C") for event in child.spec_json["_extraction_review"]["events"])
        assert not errors
        context.close()
        browser.close()
