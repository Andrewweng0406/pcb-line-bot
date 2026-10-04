from openpyxl import load_workbook

from app.core.config import settings
from app.formal_quote_export import export_formal_quote
from app.export_excel import export_quote_excel
from tests.test_web_quote_detail import _logged_in_client


def test_exports_preserve_currency_and_cents(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "EXPORT_DIR", str(tmp_path))
    spec = {"layer": 6, "qty": 10, "area_inch": 10}
    result = {"status": "success", "total": 123.45, "unit_price": 12.345}
    path = export_formal_quote(spec, result, {"currency": "USD"})
    workbook = load_workbook(path)
    assert workbook.active["D11"].value == "USD 12.35 / pc"
    assert workbook.active["E11"].value == "USD 123.45"
    workbook.close()
    filename = export_quote_excel(spec, result, currency="USD")
    workbook = load_workbook(tmp_path / filename)
    assert workbook.active["B6"].value == "USD"
    workbook.close()


def test_mixed_currency_web_and_api_do_not_show_combined_totals(temp_db):
    with temp_db.SessionLocal() as session:
        customer = temp_db.Customer(company_name="Global Buyer")
        session.add(customer)
        session.commit()
        customer_id = customer.id
    ids = []
    for currency, total in [("USD", 123.45), ("NTD", 3000)]:
        temp_db.save_quote("web:1", {"layer": 6, "qty": 1, "area_inch": 10},
                           {"status": "success", "total": total, "unit_price": total}, customer_id=customer_id)
        with temp_db.SessionLocal() as session:
            quote = session.query(temp_db.QuoteHistory).order_by(temp_db.QuoteHistory.id.desc()).first()
            quote.currency = currency
            quote.quote_outcome = "won"
            quote.final_price = total
            ids.append(quote.id)
            session.commit()
    client = _logged_in_client(temp_db)
    for route in ["/", "/quotes", "/customers", "/stats", f"/quotes/{ids[0]}"]:
        page = client.get(route)
        assert page.status_code == 200
        assert "USD 123.45" in page.text
        assert "NT$" not in page.text
        assert "3,123.45" not in page.text
    stats = client.get("/api/stats/summary").json()
    assert stats["total_amount"] is None
    assert stats["avg_price"] is None
    assert {a["currency"]: a["total"] for a in stats["amounts_by_currency"]} == {"USD": 123.45, "NTD": 3000}
    assert temp_db.get_system_stats()["avg_price"] is None
