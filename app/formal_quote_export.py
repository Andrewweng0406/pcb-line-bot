from datetime import datetime
import os
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from app.core.config import settings


def _fmt_num(value, digits=2):
    if value is None or value == "":
        return "-"
    try:
        number = float(value)
    except (TypeError, ValueError):
        return value
    if number.is_integer():
        return str(int(number))
    return f"{number:.{digits}f}".rstrip("0").rstrip(".")


def _fmt_money(value, currency):
    if value is None:
        return "-"
    return f"{(currency or '').strip().upper() or 'UNKNOWN'} {float(value):,.2f}"


def _surface(parsed):
    finish = parsed.get("surface_finish")
    gold = parsed.get("enig_thickness_uinch") or parsed.get("gold_thickness_uin")
    if finish:
        if gold and ("gold" in str(finish).lower() or str(finish).upper() == "ENIG"):
            return f'{finish} {_fmt_num(gold)}u"'
        return finish
    if parsed.get("enig"):
        return f'ENIG {_fmt_num(gold)}u"' if gold else "ENIG"
    return "-"


def _copper(parsed):
    outer = parsed.get("copper_outer_oz")
    inner = parsed.get("copper_inner_oz")
    if outer and inner:
        return f"Outer {_fmt_num(outer)} oz / Inner {_fmt_num(inner)} oz"
    return parsed.get("copper_weight") or "-"


def _special_comment(parsed):
    comments = []
    for label, key in [
        ("VIP", "vip"),
        ("Back Drill", "back_drill"),
        ("BVH", "bvh"),
        ("Countersunk", "countersunk"),
        ("Counterbored", "counterbored"),
    ]:
        if parsed.get(key):
            comments.append(label)
    if parsed.get("impedance"):
        comments.append("Impedance")
    if parsed.get("warpage_mil_per_inch"):
        comments.append(f'Board Warpage: {_fmt_num(parsed.get("warpage_mil_per_inch"))} mil/inch')
    if parsed.get("solder_mask_color"):
        comments.append(f'Solder mask: {parsed.get("solder_mask_color")}')
    if parsed.get("legend_color"):
        comments.append(f'Legend: {parsed.get("legend_color")}')
    if parsed.get("special_requirements"):
        comments.append(str(parsed.get("special_requirements")))
    return " / ".join(comments) if comments else "-"


def export_formal_quote(parsed, result, metadata=None):
    """Create a customer-facing quote without relying on a checked-in xlsx template."""
    metadata = metadata or {}

    wb = Workbook()
    ws = wb.active
    ws.title = "Customer Quote"

    blue = "1F4E8C"
    light_blue = "D9EAF7"
    gray = "EDEDED"
    border = Border(
        left=Side(style="thin", color="808080"),
        right=Side(style="thin", color="808080"),
        top=Side(style="thin", color="808080"),
        bottom=Side(style="thin", color="808080"),
    )

    ws.merge_cells("A1:H2")
    ws["A1"] = metadata.get("company_name") or "Your PCB Company"
    ws["A1"].font = Font(size=22, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=blue)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A3:H3")
    ws["A3"] = "PCB Manufacturing Services - Official Quotation"
    ws["A3"].font = Font(size=14, bold=True)
    ws["A3"].alignment = Alignment(horizontal="center")

    info_rows = [
        ("To", metadata.get("customer_name") or parsed.get("company_name") or "-"),
        ("From", metadata.get("sales_name") or "Your PCB Company - Sales"),
        ("Date", metadata.get("quote_date") or datetime.now().strftime("%Y/%m/%d")),
        ("Quote No", metadata.get("quote_no") or "-"),
    ]
    for row_index, (label, value) in enumerate(info_rows, start=5):
        ws[f"A{row_index}"] = label
        ws[f"B{row_index}"] = value
        ws[f"A{row_index}"].font = Font(bold=True)

    headers = ["Name / Part No", "Specification", "Qty", "Unit Price", "Total", "Lead Time"]
    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=10, column=col_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=gray)
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    part_name = parsed.get("pcb_name") or parsed.get("gerber_name") or parsed.get("part_no") or "-"
    lead_time = parsed.get("delivery_days") or result.get("delivery_days")
    lead_time_text = f"{lead_time} working days" if lead_time else "TBD working days"
    unit_price = result.get("unit_price")
    total = result.get("total")

    spec_lines = [
        f'Layer: {_fmt_num(parsed.get("layer"))}',
        (
            f'Single size: {_fmt_num(parsed.get("length_mm"))} x '
            f'{_fmt_num(parsed.get("width_mm"))} mm'
            if parsed.get("length_mm") and parsed.get("width_mm")
            else f'Area: {_fmt_num(result.get("area_inch") or parsed.get("area_inch"))} in2'
        ),
        f'Base Material: {parsed.get("material") or "-"}',
        f'Board thickness: {_fmt_num(parsed.get("thickness_mm") or parsed.get("thickness"))} mm',
        f'Copper thickness: {_copper(parsed)}',
        f"Surface: {_surface(parsed)}",
    ]
    optional_specs = [
        ("Min Pitch", parsed.get("pitch_mm"), "mm"),
        ("Min Hole", parsed.get("min_hole_mil"), "mil"),
        ("Line/Space", parsed.get("line_space_mil"), "mil"),
        ("Hole/Land", parsed.get("hole_land_mil"), "mil"),
        ("Aspect Ratio", parsed.get("aspect_ratio"), ""),
    ]
    for label, value, unit in optional_specs:
        if value:
            spec_lines.append(f"{label}: {_fmt_num(value)} {unit}".rstrip())
    spec_lines.append(f"Special comment: {_special_comment(parsed)}")

    row = 11
    values = [
        part_name,
        "\n".join(spec_lines),
        f'{_fmt_num(parsed.get("qty"))} pcs/batch',
        f'{_fmt_money(unit_price, metadata.get("currency", settings.DEFAULT_CURRENCY))} / pc',
        _fmt_money(total, metadata.get("currency", settings.DEFAULT_CURRENCY)),
        lead_time_text,
    ]
    for col_index, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col_index, value=value)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[row].height = 170

    ws.merge_cells("A14:H14")
    ws["A14"] = "NOTE"
    ws["A14"].font = Font(bold=True)
    ws["A14"].fill = PatternFill("solid", fgColor=light_blue)

    notes = [
        "1. Above rate quote is effective upon your confirmation of specification.",
        "2. Shipping days are not included.",
        "3. Validity: 30 days from the date of this quotation.",
        "4. Payment term: NET 30 days unless otherwise agreed.",
    ]
    for row_index, note in enumerate(notes, start=15):
        ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=8)
        ws.cell(row=row_index, column=1, value=note)

    widths = {
        "A": 26,
        "B": 52,
        "C": 16,
        "D": 18,
        "E": 18,
        "F": 18,
        "G": 4,
        "H": 4,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    exports_dir = settings.EXPORT_DIR
    os.makedirs(exports_dir, exist_ok=True)
    filename = f'formal_quote_{datetime.now().strftime("%Y%m%d_%H%M%S")}_{uuid4().hex}.xlsx'
    output_path = os.path.join(exports_dir, filename)
    wb.save(output_path)

    return output_path
