"""Safe Excel historical quote import with mapping, preview, and confirm."""

import json
import re
import hashlib
import zipfile
from datetime import datetime
from io import BytesIO
from typing import Dict, Iterable, Optional

from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError
from app.core.config import settings

from app.quote_metrics import to_non_negative_float, to_non_negative_int, calculate_margin
from app.quote_outcomes import normalize_outcome
from app.rfq_normalization import normalize_layer, normalize_material
from app.rfq_normalization import normalize_surface_finish

OPTIONAL_MAPPING = {
    "currency": "Currency", "pricing_version": "Pricing Version",
    "thickness_mm": "Thickness (mm)", "copper_weight_oz": "Copper (oz)",
    "surface_finish": "Surface Finish", "gold_thickness_uin": "Gold (uin)",
    "delivery_days": "Lead Time (days)", "final_price": "Accepted Price",
    "actual_cost": "Actual Cost", "enig": "ENIG", "vip": "VIP",
    "impedance": "Impedance", "back_drill": "Back Drill", "bvh": "BVH",
}


DEFAULT_MAPPING = {
    "customer": "Customer Name",
    "layer": "Layers",
    "material": "Material",
    "qty": "Qty",
    "size": "Size",
    "total": "Quote",
    "quote_date": "Date",
    "quote_outcome": "Result",
}


def parse_mapping_json(raw: str = "") -> dict:
    if not raw:
        return {}
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        raise ValueError("Invalid column mapping JSON")
    if not isinstance(data, dict):
        raise ValueError("Column mapping must be a JSON object")
    return {str(k): str(v) for k, v in data.items() if v}


def load_excel_rows(file_bytes: bytes) -> tuple[list[str], list[dict]]:
    if len(file_bytes) > settings.MAX_UPLOAD_SIZE:
        raise ValueError("Workbook exceeds the upload size limit")
    with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
            raise ValueError("Workbook exceeds the uncompressed size limit")
    workbook = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    try:
        worksheet = workbook.active
        if worksheet.max_column > 100 or worksheet.max_row > 5001:
            raise ValueError("Workbook limit: 5,000 data rows and 100 columns")
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return [], []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    populated = [header for header in headers if header]
    if len(populated) != len(set(populated)):
        raise ValueError("Duplicate column headers are not allowed")
    records = []
    for index, row in enumerate(rows[1:], start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        records.append({
            "row_number": index,
            "data": {
                headers[col_index]: row[col_index]
                for col_index in range(min(len(headers), len(row)))
                if headers[col_index]
            },
        })
    return headers, records


def _field(row: dict, mapping: dict, field: str):
    column = mapping.get(field)
    if not column:
        return None
    return row.get(column)


def parse_size(value) -> tuple[Optional[float], Optional[float]]:
    if value is None:
        return None, None
    text = str(value).lower().replace("mm", "")
    match = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*[x×]\s*(\d+(?:\.\d+)?)\s*", text)
    if not match:
        return None, None
    return float(match.group(1)), float(match.group(2))


def parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def normalize_result(value):
    if value is None or value == "":
        return "pending"
    text = str(value).strip().lower()
    aliases = {
        "win": "won",
        "won": "won",
        "order": "won",
        "ordered": "won",
        "lost": "lost",
        "loss": "lost",
        "no response": "no_response",
        "no_response": "no_response",
        "cancelled": "cancelled",
        "canceled": "cancelled",
        "pending": "pending",
    }
    return normalize_outcome(aliases.get(text, text))


def validate_import_rows(records: Iterable[dict], mapping: dict) -> list[dict]:
    mapping = {**OPTIONAL_MAPPING, **mapping}
    validated = []
    for record in records:
        raw = record["data"]
        errors = []
        warnings = []

        layer = normalize_layer(_field(raw, mapping, "layer"))
        qty = to_non_negative_int(_field(raw, mapping, "qty"))
        total = to_non_negative_float(_field(raw, mapping, "total"))
        material = normalize_material(_field(raw, mapping, "material"))
        length_mm = to_non_negative_float(_field(raw, mapping, "length_mm"))
        width_mm = to_non_negative_float(_field(raw, mapping, "width_mm"))
        if (length_mm is None or width_mm is None) and mapping.get("size"):
            length_mm, width_mm = parse_size(_field(raw, mapping, "size"))
        quote_date = parse_date(_field(raw, mapping, "quote_date"))
        outcome = normalize_result(_field(raw, mapping, "quote_outcome"))
        if outcome is None:
            errors.append("Invalid quote outcome")

        if layer is None or layer <= 0 or isinstance(layer, bool):
            errors.append("Missing or invalid layer")
        if qty is None or qty == 0:
            errors.append("Missing or invalid quantity")
        if total is None:
            errors.append("Missing or invalid quote total")
        if material is None:
            warnings.append("Missing material")
        if length_mm is None or width_mm is None:
            warnings.append("Missing board size")
        if any(value is not None and value <= 0 for value in (length_mm, width_mm)):
            errors.append("Board dimensions must be positive")
        if mapping.get("quote_date") and quote_date is None:
            warnings.append("Could not parse quote date")

        optional = {}
        for field in ("thickness_mm", "copper_weight_oz", "gold_thickness_uin", "delivery_days", "final_price", "actual_cost"):
            raw_value = _field(raw, mapping, field)
            parser = to_non_negative_int if field == "delivery_days" else to_non_negative_float
            value = parser(raw_value)
            if raw_value not in (None, "") and (value is None or (field in {"thickness_mm", "copper_weight_oz", "delivery_days"} and value <= 0)):
                errors.append(f"Invalid {field}")
            optional[field] = value
        for field in ("currency", "pricing_version", "surface_finish"):
            value = _field(raw, mapping, field)
            optional[field] = str(value).strip() if value is not None and str(value).strip() else None
        if optional["currency"]:
            optional["currency"] = optional["currency"].upper()
            if not re.fullmatch(r"[A-Z]{3}", optional["currency"]):
                errors.append("Currency must be a three-letter code")
        for field in ("enig", "vip", "impedance", "back_drill", "bvh"):
            raw_value = _field(raw, mapping, field)
            value = None
            if raw_value not in (None, ""):
                normalized = str(raw_value).strip().lower()
                if normalized in {"true", "yes", "1"}:
                    value = True
                elif normalized in {"false", "no", "0"}:
                    value = False
                else:
                    errors.append(f"Invalid {field}; use yes/no or true/false")
            optional[field] = value
        missing = [field for field in ("currency", "pricing_version", "thickness_mm", "copper_weight_oz", "surface_finish", "delivery_days", "enig", "vip", "impedance", "back_drill", "bvh") if optional[field] is None]
        optional["surface_finish"] = normalize_surface_finish(optional["surface_finish"])
        if optional["surface_finish"] in {"ENIG", "Hard Gold"} and (optional["gold_thickness_uin"] is None or optional["gold_thickness_uin"] <= 0):
            warnings.append("Reference data missing: gold_thickness_uin")
        if missing:
            warnings.append("Reference data missing: " + ", ".join(missing))

        validated.append({
            "row_number": record["row_number"],
            "status": "invalid" if errors else "valid",
            "errors": errors,
            "warnings": warnings,
            "mapped": {
                **optional,
                "customer": _field(raw, mapping, "customer"),
                "layer": layer,
                "material": material,
                "qty": qty,
                "length_mm": length_mm,
                "width_mm": width_mm,
                "total": total,
                "unit_price": round(total / qty, 2) if total is not None and qty else None,
                "quote_date": quote_date,
                "quote_outcome": outcome,
            },
        })
    return validated


def preview_import(file_bytes: bytes, mapping: dict = None) -> dict:
    headers, records = load_excel_rows(file_bytes)
    mapping = mapping or DEFAULT_MAPPING
    rows = validate_import_rows(records, mapping) if mapping else []
    valid = [row for row in rows if row["status"] == "valid"]
    invalid = [row for row in rows if row["status"] == "invalid"]
    warnings = [row for row in rows if row["warnings"]]
    return {
        "headers": headers,
        "rows_detected": len(records),
        "valid": len(valid),
        "warnings": len(warnings),
        "invalid": len(invalid),
        "rows": rows,
        "preview_rows": rows[:50],
        "mapping": {**OPTIONAL_MAPPING, **mapping},
    }


def confirm_import(session, db_module, file_bytes: bytes, mapping: dict = None, user_id: int = None) -> dict:
    preview = preview_import(file_bytes, mapping or DEFAULT_MAPPING)
    valid_rows = [row for row in preview["rows"] if row["status"] == "valid"]
    if preview["invalid"]:
        return {**preview, "committed": 0, "status": "error"}

    committed = 0
    duplicates = 0
    file_hash = hashlib.sha256(file_bytes).hexdigest()
    mapping_hash = hashlib.sha256(json.dumps({**OPTIONAL_MAPPING, **(mapping or DEFAULT_MAPPING)}, sort_keys=True).encode()).hexdigest()
    for row in valid_rows:
        mapped = row["mapped"]
        import_key = hashlib.sha256(f"{file_hash}:{row['row_number']}".encode()).hexdigest()
        existing = session.query(db_module.QuoteHistory).filter(db_module.QuoteHistory.import_key == import_key).first()
        if existing:
            if (existing.spec_json or {}).get("mapping_sha256") != mapping_hash:
                session.rollback()
                raise ValueError("This row was already imported with a different mapping. Revise the existing quote instead.")
            duplicates += 1
            continue
        customer_id = None
        customer_name = mapped.get("customer")
        if customer_name:
            customer_name = str(customer_name).strip()
            customer = (
                session.query(db_module.Customer)
                .filter(db_module.Customer.company_name == customer_name)
                .first()
            )
            if customer is None:
                customer = db_module.Customer(company_name=customer_name)
                session.add(customer)
                session.flush()
            customer_id = customer.id

        quote = db_module.QuoteHistory(
            import_key=import_key,
            quote_no=db_module._generate_quote_no(session),
            source_channel_id="excel_import",
            source_channel="excel_import",
            product_type="pcb",
            customer_id=customer_id,
            layer=mapped["layer"],
            material=mapped["material"],
            length_mm=mapped["length_mm"],
            width_mm=mapped["width_mm"],
            qty=mapped["qty"],
            total=mapped["total"],
            unit_price=mapped["unit_price"],
            status="pending",
            quote_outcome=mapped["quote_outcome"],
            currency=mapped["currency"],
            pricing_version=mapped["pricing_version"],
            board_thickness_mm=mapped["thickness_mm"],
            copper_weight_oz=mapped["copper_weight_oz"],
            surface_finish=mapped["surface_finish"],
            gold_thickness_uin=mapped["gold_thickness_uin"],
            delivery_days=mapped["delivery_days"],
            final_price=mapped["final_price"],
            actual_cost=mapped["actual_cost"],
            actual_margin_pct=calculate_margin(mapped["final_price"], mapped["actual_cost"]),
            rfq_received_at=mapped["quote_date"],
            quote_sent_at=mapped["quote_date"],
            spec_json={
                "source": "excel_import",
                "file_sha256": file_hash,
                "mapping_sha256": mapping_hash,
                **{field: mapped[field] for field in OPTIONAL_MAPPING},
                "row_number": row["row_number"],
                "layer": mapped["layer"],
                "material": mapped["material"],
                "qty": mapped["qty"],
                "length_mm": mapped["length_mm"],
                "width_mm": mapped["width_mm"],
            },
            breakdown_json={"source": "excel_import"},
            created_by_user_id=user_id,
        )
        session.add(quote)
        try:
            session.flush()
        except IntegrityError:
            session.rollback()
            raise ValueError("Import conflicted with another request. Retry; no rows from this request were committed.")
        committed += 1
    try:
        session.commit()
    except IntegrityError:
        session.rollback()
        raise ValueError("Import conflicted with another request. Preview and retry; no rows from this request were committed.")
    return {**preview, "committed": committed, "duplicates": duplicates, "status": "success"}
